from pathlib import Path
from typing import Dict, List, Optional, Union
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from src.logger import logger

class ExcelWriter:
    """Formats and writes master datasets and error reports to styled .xlsx workbooks."""

    HEADER_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep navy
    HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    DATA_FONT = Font(name="Segoe UI", size=10)
    ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Light slate
    
    THIN_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    @classmethod
    def write_dataframe(
        cls,
        df: pd.DataFrame,
        sheet_title: str = "Sheet1",
        output_path: Optional[Union[str, Path]] = None,
        header_color: str = "1E3A8A"
    ) -> Union[Path, io.BytesIO]:
        """
        Exports a DataFrame to an aesthetically styled Excel workbook.
        Returns Path if output_path is provided, or io.BytesIO otherwise.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title

        # Drop internal tracking columns if present
        clean_df = df.copy()
        for internal_col in ["_excel_row_num", "_source_file", "_duplicate_key"]:
            if internal_col in clean_df.columns and sheet_title != "Error Report":
                clean_df.drop(columns=[internal_col], inplace=True)

        # Write headers
        headers = list(clean_df.columns)
        ws.append(headers)

        custom_header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")

        # Style header row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = custom_header_fill
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cls.THIN_BORDER

        ws.row_dimensions[1].height = 26

        # Write rows
        for row_idx, row_values in enumerate(clean_df.itertuples(index=False), start=2):
            ws.append(list(row_values))
            ws.row_dimensions[row_idx].height = 20
            is_alt = (row_idx % 2 == 0)
            
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = cls.DATA_FONT
                cell.border = cls.THIN_BORDER
                if is_alt:
                    cell.fill = cls.ALT_FILL
                
                # Basic alignment heuristic
                val = cell.value
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        for col_idx, col in enumerate(clean_df.columns, start=1):
            col_letter = get_column_letter(col_idx)
            max_len = len(str(col))
            for row in range(2, min(ws.max_row + 1, 100)): # sample first 100 rows for speed
                cell_val = str(ws.cell(row=row, column=col_idx).value or "")
                max_len = max(max_len, len(cell_val))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

        # Freeze top row
        ws.freeze_panes = "A2"

        if output_path:
            out_p = Path(output_path)
            out_p.parent.mkdir(parents=True, exist_ok=True)
            wb.save(out_p)
            logger.info(f"Saved Excel file to: {out_p}")
            return out_p
        else:
            stream = io.BytesIO()
            wb.save(stream)
            stream.seek(0)
            return stream

    @classmethod
    def write_error_report(
        cls,
        errors_df: pd.DataFrame,
        output_path: Optional[Union[str, Path]] = None
    ) -> Union[Path, io.BytesIO]:
        """Writes the error report with distinctive warning/alert styling (dark crimson header)."""
        return cls.write_dataframe(
            df=errors_df,
            sheet_title="Error Report",
            output_path=output_path,
            header_color="991B1B" # Crimson red
        )
